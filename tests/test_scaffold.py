import datetime
import io
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from contextlib import redirect_stdout
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SCRIPT_DIR = ROOT / "scripts" / "split_pattern"
sys.path.insert(0, str(SCRIPT_DIR))

import create_thesis_docx as generator  # noqa: E402


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CP_NS = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
DC_NS = "http://purl.org/dc/elements/1.1/"
DCTERMS_NS = "http://purl.org/dc/terms/"


def w_tag(local_name):
    return f"{{{W_NS}}}{local_name}"


class ScaffoldGenerationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.output_path = Path(cls.temp_dir.name) / "scaffold.docx"

        original_output_path = generator.OUTPUT_PATH
        generator.OUTPUT_PATH = str(cls.output_path)
        stdout = io.StringIO()
        try:
            with redirect_stdout(stdout):
                cls.result_path = generator.create_document()
        finally:
            generator.OUTPUT_PATH = original_output_path
        cls.stdout = stdout.getvalue()

    @classmethod
    def tearDownClass(cls):
        cls.temp_dir.cleanup()

    def test_missing_figures_are_explicit_placeholders(self):
        self.assertEqual(self.result_path, str(self.output_path))
        self.assertTrue(self.output_path.is_file())
        self.assertEqual(
            self.stdout.count("[警告] 图片未找到"),
            len(generator.FIGURE_MAP),
        )

        document = Document(self.output_path)
        placeholders = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.startswith("[图片待插入:")
        ]
        self.assertEqual(len(placeholders), len(generator.FIGURE_MAP))

    def test_front_matter_and_body_use_separate_page_numbering(self):
        # Suzhou University's public 2020 print format requires lowercase Roman
        # front-matter pages and a decimal body section restarting at 1.
        document = Document(self.output_path)
        self.assertEqual(len(document.sections), 2)
        self.assertEqual(document.sections[0].footer.paragraphs[0].text, "")
        self.assertEqual(document.sections[1].footer.paragraphs[0].text, "--")

        with zipfile.ZipFile(self.output_path) as archive:
            self.assertIsNone(archive.testzip())
            document_xml = ET.fromstring(archive.read("word/document.xml"))
            settings_xml = ET.fromstring(archive.read("word/settings.xml"))

        section_properties = list(document_xml.iter(w_tag("sectPr")))
        self.assertEqual(len(section_properties), 2)
        page_numbering = [
            section.find(w_tag("pgNumType"))
            for section in section_properties
        ]
        self.assertEqual(
            [node.get(w_tag("fmt")) for node in page_numbering],
            ["lowerRoman", "decimal"],
        )
        self.assertEqual(
            [node.get(w_tag("start")) for node in page_numbering],
            ["1", "1"],
        )

        update_fields = settings_xml.find(w_tag("updateFields"))
        self.assertIsNotNone(update_fields)
        self.assertEqual(update_fields.get(w_tag("val")), "true")

        field_begins = [
            element
            for element in document_xml.iter(w_tag("fldChar"))
            if element.get(w_tag("fldCharType")) == "begin"
        ]
        self.assertEqual(len(field_begins), len(generator.TOC_ENTRIES) + 1)
        self.assertTrue(all(node.get(w_tag("dirty")) == "true" for node in field_begins))

    def test_generated_ooxml_inserts_schema_sensitive_children_in_order(self):
        with zipfile.ZipFile(self.output_path) as archive:
            document_xml = ET.fromstring(archive.read("word/document.xml"))
            settings_xml = ET.fromstring(archive.read("word/settings.xml"))

        section_properties = list(document_xml.iter(w_tag("sectPr")))
        self.assertEqual(len(section_properties), 2)
        for section in section_properties:
            child_tags = [child.tag for child in section]
            self.assertLess(
                child_tags.index(w_tag("pgMar")),
                child_tags.index(w_tag("pgNumType")),
            )
            self.assertLess(
                child_tags.index(w_tag("pgNumType")),
                child_tags.index(w_tag("cols")),
            )

        settings_child_tags = [child.tag for child in settings_xml]
        self.assertLess(
            settings_child_tags.index(w_tag("updateFields")),
            settings_child_tags.index(w_tag("compat")),
        )

    def test_cached_toc_pages_match_expected_scaffold_pagination(self):
        # This checks the deliberately cached values only. Pagination/render QA is a
        # separate release check because this unit test does not render the DOCX.
        expected_pages = [
            1, 1, 1, 1, 2,
            3, 3, 3, 3, 3, 4,
            5, 5, 5, 6, 6, 6, 7, 7,
            8, 8, 8, 8, 8, 9,
            10, 11, 12,
        ]
        self.assertEqual(
            [page for _, page, _ in generator.TOC_ENTRIES],
            expected_pages,
        )

    def test_personal_metadata_is_cleared_and_timestamps_are_utc(self):
        document = Document(self.output_path)
        self.assertEqual(document.core_properties.author, "")
        self.assertEqual(document.core_properties.last_modified_by, "")

        with zipfile.ZipFile(self.output_path) as archive:
            names = archive.namelist()
            self.assertFalse(any(name.startswith("word/comments") for name in names))
            core_xml = ET.fromstring(archive.read("docProps/core.xml"))

        self.assertEqual(core_xml.find(f"{{{DC_NS}}}creator").text, None)
        self.assertEqual(core_xml.find(f"{{{CP_NS}}}lastModifiedBy").text, None)

        created_text = core_xml.find(f"{{{DCTERMS_NS}}}created").text
        created = datetime.datetime.fromisoformat(created_text.replace("Z", "+00:00"))
        age = datetime.datetime.now(datetime.timezone.utc) - created
        self.assertGreaterEqual(age.total_seconds(), 0)
        self.assertLess(age.total_seconds(), 120)


if __name__ == "__main__":
    unittest.main()
